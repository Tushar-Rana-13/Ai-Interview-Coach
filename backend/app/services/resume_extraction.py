
from pathlib import Path

import pymupdf
from docx import Document
from fastapi import HTTPException, status


def extract_resume_text(file_path: str) -> str:
    """Extract readable text from a stored PDF or DOCX resume."""

    path = Path(file_path)

    if not path.is_file():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Stored resume file not found.",
        )

    extension = path.suffix.lower()

    try:
        if extension == ".pdf":
            with pymupdf.open(path) as document:
                if document.is_encrypted:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                        detail="Password-protected PDFs are not supported.",
                    )

                text = "\n".join(
                    page.get_text(sort=True) for page in document
                )

        elif extension == ".docx":
            document = Document(path)

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            table_rows = [
                " | ".join(
                    cell.text.strip() for cell in row.cells
                )
                for table in document.tables
                for row in table.rows
            ]

            text = "\n".join(paragraphs + table_rows)

        else:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail="Only PDF and DOCX files are supported.",
            )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The resume could not be read. Please upload a valid document.",
        ) from exc

    cleaned_text = "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )

    if not cleaned_text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "No readable text was found. "
                "The document may be scanned or image-based."
            ),
        )

    return cleaned_text
